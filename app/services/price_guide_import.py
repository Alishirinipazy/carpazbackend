"""
آپلود/دانلود اکسل برای «قیمت روز خودرو».

فرمت مورد انتظار فایل اکسل (ردیف اول = هدر، دقیقاً به همین ترتیب ستون‌ها):
برند | مدل | تیریم | سال | وضعیت (صفر/کارکرده) | قیمت بازار | قیمت نمایندگی

- ستون‌های «تیریم»، «سال» و «قیمت نمایندگی» اختیاری‌اند و می‌توانند خالی باشند.
- مقدار ستون «وضعیت» باید «صفر» یا «کارکرده» باشد (بدون حساسیت به فاصله).
- قیمت‌ها به تومان و فقط عدد (بدون کاما/متن) نوشته شوند؛ اگر کاما داشته باشند
  هم پاک‌سازی می‌شوند.
"""
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.models.price_guide import CONDITION_USED, CONDITION_ZERO

EXPECTED_HEADERS = ["برند", "مدل", "تیریم", "سال", "وضعیت", "قیمت بازار", "قیمت نمایندگی"]


def _clean_price(value) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value)
    cleaned = str(value).replace(",", "").replace("،", "").strip()
    if not cleaned:
        return None
    try:
        return int(float(cleaned))
    except ValueError:
        return None


def _parse_condition(value) -> int | None:
    if value is None:
        return None
    v = str(value).strip()
    if v in ("صفر", "0", "zero", "صفر کیلومتر"):
        return CONDITION_ZERO
    if v in ("کارکرده", "1", "used"):
        return CONDITION_USED
    return None


def parse_price_guide_excel(file_bytes: bytes) -> tuple[list[dict], list[str]]:
    """
    فایل اکسل را می‌خواند و برمی‌گرداند: (لیست ردیف‌های معتبر، لیست خطاهای هر ردیف نامعتبر).
    هر ردیف معتبر یک dict متناظر با فیلدهای PriceGuideEntryIn است.
    """
    wb = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = wb.active

    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None)

    valid_rows: list[dict] = []
    errors: list[str] = []

    for i, row in enumerate(rows_iter, start=2):  # ردیف ۱ هدر است
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue  # ردیف خالی را نادیده بگیر

        try:
            brand = str(row[0] or "").strip()
            model = str(row[1] or "").strip()
            trim = str(row[2]).strip() if len(row) > 2 and row[2] not in (None, "") else None
            year = str(row[3]).strip() if len(row) > 3 and row[3] not in (None, "") else None
            condition = _parse_condition(row[4] if len(row) > 4 else None)
            market_price = _clean_price(row[5] if len(row) > 5 else None)
            agency_price = _clean_price(row[6] if len(row) > 6 else None)

            if not brand or not model:
                errors.append(f"ردیف {i}: برند و مدل الزامی است")
                continue
            if condition is None:
                errors.append(f"ردیف {i}: ستون وضعیت باید «صفر» یا «کارکرده» باشد")
                continue
            if market_price is None and agency_price is None:
                errors.append(f"ردیف {i}: حداقل یکی از قیمت بازار یا نمایندگی باید مقدار داشته باشد")
                continue

            valid_rows.append({
                "brand": brand,
                "model": model,
                "trim": trim,
                "year": year,
                "condition": condition,
                "market_price": market_price,
                "agency_price": agency_price,
            })
        except Exception as e:  # noqa: BLE001 - یک ردیف خراب نباید کل فایل را متوقف کند
            errors.append(f"ردیف {i}: خطای غیرمنتظره ({e})")

    return valid_rows, errors


def build_price_guide_excel_template() -> bytes:
    """یک فایل نمونه/قالب اکسل برای دانلود توسط ادمین می‌سازد."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = "قیمت روز خودرو"
    sheet.append(EXPECTED_HEADERS)
    sheet.append(["ایران خودرو", "دنا پلاس", "توربو", "1403", "صفر", "1850000000", "1650000000"])
    sheet.append(["پژو", "207", "پانوراما اتوماتیک", "1401", "کارکرده", "980000000", ""])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
